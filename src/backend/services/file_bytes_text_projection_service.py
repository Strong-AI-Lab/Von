"""Bounded, effect-free text projection for file bytes.

This module deliberately owns no persistence, indexing, enrichment, or model
calls.  It is suitable for inspecting private source bytes without first
turning them into a durable ``computer_file_copy``.
"""

from __future__ import annotations

import io
import json
import os
import re
import signal
import struct
import subprocess
import sys
import threading
import zipfile
from collections.abc import Mapping
from email import policy as email_policy
from email.parser import BytesParser
from html import unescape
from pathlib import Path
from typing import Any

DEFAULT_MAX_TEXT_CHARS = 20_000
MAX_TEXT_CHARS = 250_000

# These are parser resource boundaries, not turn deadlines.  The input ceiling
# matches the Gmail attachment read boundary; the narrower format ceilings
# prevent parsers from multiplying a bounded source into unbounded in-process
# work.  They intentionally remain local to this effect-free projection.
MAX_INPUT_BYTES = 50_000_000
MAX_MARKUP_OR_EMAIL_BYTES = 10_000_000
MAX_ARCHIVE_MEMBERS = 2_000
MAX_ARCHIVE_EXPANDED_BYTES = 100_000_000
MAX_ARCHIVE_MEMBER_BYTES = 25_000_000
MAX_ARCHIVE_XML_ELEMENTS = 500_000
MAX_PDF_PAGES = 500
PDF_PARSER_MEMORY_LIMIT_BYTES = 1_073_741_824
PDF_PARSER_CPU_LIMIT_SECONDS = 5
PDF_PARSER_WALL_LIMIT_SECONDS = 10.0
PDF_PARSER_PROTOCOL_MAX_BYTES = 2_000_000
MAX_DOCX_PARAGRAPHS = 20_000
MAX_DOCX_TABLES = 1_000
MAX_DOCX_ROWS = 20_000
MAX_DOCX_CELLS = 100_000
MAX_PPTX_SLIDES = 500
MAX_PPTX_SHAPES = 20_000
MAX_XLSX_WORKSHEETS = 200
MAX_XLSX_ROWS = 10_000
MAX_XLSX_COLUMNS = 256
MAX_XLSX_CELLS = 100_000
MAX_XLSX_SHARED_STRINGS = 100_000
MAX_XLSX_STYLE_RECORDS = 20_000

_ZIP_END_OF_CENTRAL_DIRECTORY = b"PK\x05\x06"
_ZIP_END_OF_CENTRAL_DIRECTORY_STRUCT = struct.Struct("<4s4H2LH")
_ZIP_CENTRAL_DIRECTORY = b"PK\x01\x02"
_ZIP_CENTRAL_DIRECTORY_STRUCT = struct.Struct("<4s6H3L5H2L")
_SAFE_ZIP_COMPRESSION_TYPES = {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
_XML_START_TAG = re.compile(
    rb"<(?:[A-Za-z_][A-Za-z0-9_.-]*:)?([A-Za-z_][A-Za-z0-9_.-]*)(?=[\s/>])"
)
_PPTX_SHAPE_TAGS = {b"sp", b"pic", b"graphicFrame", b"grpSp", b"cxnSp"}


class _ProjectionResourceLimit(Exception):
    """A named local parser boundary was reached before unsafe expansion."""


def _append_bounded_text(
    parts: list[str],
    value: str,
    *,
    length: int,
    stop_after_chars: int,
) -> tuple[int, bool]:
    """Append no more projection text than is needed to prove truncation."""

    remaining = max(0, stop_after_chars - length)
    if remaining:
        parts.append(value[:remaining])
    new_length = length + len(value)
    return min(new_length, stop_after_chars), new_length >= stop_after_chars


def _zip_directory_location_before_open(data: bytes) -> tuple[int, int, int]:
    """Read the non-ZIP64 member count without materialising ZipInfo objects."""

    search_from = max(0, len(data) - (65_535 + 22))
    offset = data.rfind(_ZIP_END_OF_CENTRAL_DIRECTORY, search_from)
    record: tuple[int, int, int] | None = None
    while offset >= 0:
        if offset + _ZIP_END_OF_CENTRAL_DIRECTORY_STRUCT.size <= len(data):
            unpacked = _ZIP_END_OF_CENTRAL_DIRECTORY_STRUCT.unpack_from(data, offset)
            comment_length = unpacked[-1]
            if (
                offset + _ZIP_END_OF_CENTRAL_DIRECTORY_STRUCT.size + comment_length
                == len(data)
            ):
                (
                    _signature,
                    disk_number,
                    central_directory_disk,
                    entries_on_disk,
                    total_entries,
                    central_directory_size,
                    central_directory_offset,
                    _comment_length,
                ) = unpacked
                if (
                    disk_number
                    or central_directory_disk
                    or entries_on_disk != total_entries
                ):
                    raise _ProjectionResourceLimit("archive_multidisk_not_supported")
                if total_entries == 0xFFFF:
                    raise _ProjectionResourceLimit("archive_zip64_not_supported")
                if total_entries > MAX_ARCHIVE_MEMBERS:
                    raise _ProjectionResourceLimit("archive_member_limit_exceeded")
                if central_directory_offset + central_directory_size != offset:
                    raise zipfile.BadZipFile("invalid central directory location")
                record = (
                    total_entries,
                    central_directory_offset,
                    central_directory_size,
                )
                break
        offset = data.rfind(_ZIP_END_OF_CENTRAL_DIRECTORY, search_from, offset)
    if record is None:
        raise zipfile.BadZipFile("end of central directory not found")
    return record


def _preflight_ooxml_archive(data: bytes, *, document_kind: str) -> None:
    """Reject archive expansion beyond the local OOXML parser envelope."""

    (
        declared_member_count,
        central_directory_offset,
        central_directory_size,
    ) = _zip_directory_location_before_open(data)
    cursor = central_directory_offset
    central_directory_end = cursor + central_directory_size
    expanded_bytes = 0
    parsed_member_count = 0
    while cursor < central_directory_end:
        if cursor + _ZIP_CENTRAL_DIRECTORY_STRUCT.size > central_directory_end:
            raise zipfile.BadZipFile("truncated central directory")
        fields = _ZIP_CENTRAL_DIRECTORY_STRUCT.unpack_from(data, cursor)
        if fields[0] != _ZIP_CENTRAL_DIRECTORY:
            raise zipfile.BadZipFile("invalid central directory signature")
        flags = fields[3]
        compression_type = fields[4]
        compressed_size = fields[8]
        member_size = fields[9]
        filename_size = fields[10]
        extra_size = fields[11]
        comment_size = fields[12]
        disk_start = fields[13]
        parsed_member_count += 1
        if parsed_member_count > MAX_ARCHIVE_MEMBERS:
            raise _ProjectionResourceLimit("archive_member_limit_exceeded")
        if disk_start:
            raise _ProjectionResourceLimit("archive_multidisk_not_supported")
        if flags & 0x1:
            raise _ProjectionResourceLimit("archive_encryption_not_supported")
        if compression_type not in _SAFE_ZIP_COMPRESSION_TYPES:
            raise _ProjectionResourceLimit("archive_compression_not_supported")
        if compressed_size == 0xFFFFFFFF or member_size == 0xFFFFFFFF:
            raise _ProjectionResourceLimit("archive_zip64_not_supported")
        if member_size > MAX_ARCHIVE_MEMBER_BYTES:
            raise _ProjectionResourceLimit("archive_member_size_limit_exceeded")
        expanded_bytes += member_size
        if expanded_bytes > MAX_ARCHIVE_EXPANDED_BYTES:
            raise _ProjectionResourceLimit("archive_expansion_limit_exceeded")
        cursor += (
            _ZIP_CENTRAL_DIRECTORY_STRUCT.size
            + filename_size
            + extra_size
            + comment_size
        )
    if cursor != central_directory_end or parsed_member_count != declared_member_count:
        raise zipfile.BadZipFile("central directory member count mismatch")

    xml_element_count = 0
    document_paragraphs = 0
    document_tables = 0
    document_rows = 0
    document_cells = 0
    presentation_slides = 0
    presentation_shapes = 0
    workbook_sheets = 0
    workbook_rows = 0
    workbook_cells = 0
    workbook_shared_strings = 0
    workbook_style_records = 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) != declared_member_count:
            raise zipfile.BadZipFile("central directory member count mismatch")
        for member in members:
            member_name = member.filename.replace("\\", "/")
            if not member_name.lower().endswith((".xml", ".rels", ".vml")):
                continue
            member_xml = archive.read(member)
            is_docx_document = (
                document_kind == "docx" and member_name == "word/document.xml"
            )
            is_pptx_slide = (
                document_kind == "pptx"
                and member_name.startswith("ppt/slides/")
                and "/_rels/" not in member_name
                and member_name.endswith(".xml")
            )
            is_xlsx_sheet = (
                document_kind == "xlsx"
                and member_name.startswith("xl/worksheets/")
                and "/_rels/" not in member_name
                and member_name.endswith(".xml")
            )
            if is_pptx_slide:
                presentation_slides += 1
                if presentation_slides > MAX_PPTX_SLIDES:
                    raise _ProjectionResourceLimit("pptx_slide_limit_exceeded")
            if is_xlsx_sheet:
                workbook_sheets += 1
                if workbook_sheets > MAX_XLSX_WORKSHEETS:
                    raise _ProjectionResourceLimit("xlsx_worksheet_limit_exceeded")
            for match in _XML_START_TAG.finditer(member_xml):
                local_name = match.group(1)
                xml_element_count += 1
                if xml_element_count > MAX_ARCHIVE_XML_ELEMENTS:
                    raise _ProjectionResourceLimit("archive_xml_element_limit_exceeded")
                if is_docx_document:
                    if local_name == b"p":
                        document_paragraphs += 1
                        if document_paragraphs > MAX_DOCX_PARAGRAPHS:
                            raise _ProjectionResourceLimit(
                                "docx_paragraph_limit_exceeded"
                            )
                    elif local_name == b"tbl":
                        document_tables += 1
                        if document_tables > MAX_DOCX_TABLES:
                            raise _ProjectionResourceLimit("docx_table_limit_exceeded")
                    elif local_name == b"tr":
                        document_rows += 1
                        if document_rows > MAX_DOCX_ROWS:
                            raise _ProjectionResourceLimit("docx_row_limit_exceeded")
                    elif local_name == b"tc":
                        document_cells += 1
                        if document_cells > MAX_DOCX_CELLS:
                            raise _ProjectionResourceLimit("docx_cell_limit_exceeded")
                elif is_pptx_slide and local_name in _PPTX_SHAPE_TAGS:
                    presentation_shapes += 1
                    if presentation_shapes > MAX_PPTX_SHAPES:
                        raise _ProjectionResourceLimit("pptx_shape_limit_exceeded")
                elif is_xlsx_sheet:
                    if local_name == b"row":
                        workbook_rows += 1
                        if workbook_rows > MAX_XLSX_ROWS:
                            raise _ProjectionResourceLimit("xlsx_row_limit_exceeded")
                    elif local_name == b"c":
                        workbook_cells += 1
                        if workbook_cells > MAX_XLSX_CELLS:
                            raise _ProjectionResourceLimit("xlsx_cell_limit_exceeded")
                elif (
                    document_kind == "xlsx"
                    and member_name == "xl/sharedStrings.xml"
                    and local_name == b"si"
                ):
                    workbook_shared_strings += 1
                    if workbook_shared_strings > MAX_XLSX_SHARED_STRINGS:
                        raise _ProjectionResourceLimit(
                            "xlsx_shared_string_limit_exceeded"
                        )
                elif (
                    document_kind == "xlsx"
                    and member_name == "xl/styles.xml"
                    and local_name == b"xf"
                ):
                    workbook_style_records += 1
                    if workbook_style_records > MAX_XLSX_STYLE_RECORDS:
                        raise _ProjectionResourceLimit(
                            "xlsx_style_record_limit_exceeded"
                        )


def _clean_text(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    cleaned = value.strip()
    return cleaned or None


def _content_type_token(value: Any) -> str:
    cleaned = _clean_text(value)
    return cleaned.split(";", 1)[0].strip().lower() if cleaned else ""


def _filename_token(value: Any) -> str:
    cleaned = _clean_text(value)
    return cleaned.lower() if cleaned else ""


def _decode_text(
    data: bytes, *, stop_after_chars: int | None = None
) -> tuple[str, str, bool]:
    decode_bytes = data
    byte_truncated = False
    if stop_after_chars is not None:
        # Four source bytes per requested character covers UTF-8 and leaves a
        # small suffix for an incomplete final code point.  Text is truncated
        # again after decoding, so this bounds allocation without weakening the
        # caller's character limit.
        source_limit = (stop_after_chars * 4) + 8
        if len(decode_bytes) > source_limit:
            decode_bytes = decode_bytes[:source_limit]
            byte_truncated = True
    if decode_bytes.startswith((b"\xff\xfe", b"\xfe\xff")):
        return (
            decode_bytes.decode("utf-16", errors="replace"),
            "utf-16",
            byte_truncated,
        )
    if decode_bytes.startswith(b"\xef\xbb\xbf"):
        return (
            decode_bytes.decode("utf-8-sig", errors="replace"),
            "utf-8-sig",
            byte_truncated,
        )
    try:
        return decode_bytes.decode("utf-8"), "utf-8", byte_truncated
    except UnicodeDecodeError as exc:
        if byte_truncated and exc.reason == "unexpected end of data":
            return (
                decode_bytes.decode("utf-8", errors="replace"),
                "utf-8",
                True,
            )
        return (
            decode_bytes.decode("latin-1", errors="replace"),
            "latin-1",
            byte_truncated,
        )


def _looks_textual(data: bytes) -> bool:
    sample = data[:8_192]
    if not sample or b"\x00" in sample:
        return False
    printable = sum(
        byte in (9, 10, 13) or 32 <= byte <= 126 or byte >= 160 for byte in sample
    )
    return printable / len(sample) >= 0.85


def _pdf_parser_isolation_supported() -> bool:
    """Return whether this host has the verified hard parser-limit mechanism."""

    # Darwin advertises RLIMIT_AS but a Python process starts with a very large
    # mapped address space and cannot lower that limit to a useful parser
    # boundary.  Other platforms need their own verified confinement mechanism
    # (for example, Windows job objects) before PDF parsing may be enabled.
    if os.name != "posix" or sys.platform != "linux":
        return False
    try:
        import resource
    except ImportError:
        return False
    return bool(hasattr(resource, "RLIMIT_AS") and hasattr(resource, "RLIMIT_CPU"))


def _pdf_worker_command(*, stop_after_chars: int) -> list[str]:
    worker_path = Path(__file__).with_name("pdf_text_projection_worker.py")
    return [
        sys.executable,
        "-I",
        str(worker_path),
        "--memory-limit-bytes",
        str(PDF_PARSER_MEMORY_LIMIT_BYTES),
        "--cpu-limit-seconds",
        str(PDF_PARSER_CPU_LIMIT_SECONDS),
        "--max-input-bytes",
        str(MAX_INPUT_BYTES),
        "--max-pages",
        str(MAX_PDF_PAGES),
        "--stop-after-chars",
        str(stop_after_chars),
    ]


def _kill_pdf_worker(process: subprocess.Popen[bytes]) -> None:
    try:
        os.killpg(process.pid, signal.SIGKILL)
    except OSError:
        try:
            process.kill()
        except OSError:
            pass
    try:
        process.wait(timeout=1.0)
    except (OSError, subprocess.TimeoutExpired):
        pass


def _run_isolated_pdf_worker(
    data: bytes,
    *,
    stop_after_chars: int,
) -> dict[str, Any]:
    """Run the PDF parser with bounded input, output, memory, CPU, and wall time."""

    if not _pdf_parser_isolation_supported():
        return {
            "status": "isolation_unsupported",
            "error": "pdf_parser_isolation_unavailable",
        }

    try:
        process = subprocess.Popen(
            _pdf_worker_command(stop_after_chars=stop_after_chars),
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            close_fds=True,
            start_new_session=True,
        )
    except (OSError, ValueError):
        return {
            "status": "isolation_unsupported",
            "error": "pdf_parser_worker_start_failed",
        }

    input_stream = process.stdin
    output_stream = process.stdout
    assert input_stream is not None
    assert output_stream is not None
    output_chunks: list[bytes] = []
    output_size = 0
    output_limit_reached = threading.Event()

    def _write_input() -> None:
        try:
            input_stream.write(data)
            input_stream.close()
        except (BrokenPipeError, OSError, ValueError):
            pass

    def _read_output() -> None:
        nonlocal output_size
        try:
            while chunk := output_stream.read(65_536):
                remaining = PDF_PARSER_PROTOCOL_MAX_BYTES - output_size
                if remaining > 0:
                    output_chunks.append(chunk[:remaining])
                    output_size += min(len(chunk), remaining)
                if len(chunk) > remaining:
                    output_limit_reached.set()
                    try:
                        process.kill()
                    except OSError:
                        pass
                    break
        except (OSError, ValueError):
            pass

    writer = threading.Thread(target=_write_input, daemon=True)
    reader = threading.Thread(target=_read_output, daemon=True)
    writer.start()
    reader.start()

    wall_limit_reached = False
    try:
        return_code = process.wait(timeout=PDF_PARSER_WALL_LIMIT_SECONDS)
    except subprocess.TimeoutExpired:
        wall_limit_reached = True
        _kill_pdf_worker(process)
        return_code = process.returncode
    writer.join(timeout=1.0)
    reader.join(timeout=1.0)

    if wall_limit_reached:
        return {
            "status": "resource_limit",
            "error": "pdf_parser_wall_time_limit_exceeded",
        }
    if output_limit_reached.is_set():
        return {
            "status": "resource_limit",
            "error": "pdf_parser_protocol_output_limit_exceeded",
        }

    output = b"".join(output_chunks)
    try:
        payload = json.loads(output.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        payload = None
    if isinstance(payload, Mapping) and payload.get("schema_version") == (
        "pdf_text_projection_worker.v1"
    ):
        return dict(payload)

    if isinstance(return_code, int) and return_code < 0:
        signal_number = -return_code
        if signal_number == getattr(signal, "SIGXCPU", -1):
            error = "pdf_parser_cpu_limit_exceeded"
        else:
            error = "pdf_parser_process_resource_limit_exceeded"
        return {"status": "resource_limit", "error": error}
    return {
        "status": "worker_failure",
        "error": "pdf_parser_worker_protocol_failed",
    }


def _extract_pdf(
    data: bytes, *, stop_after_chars: int
) -> tuple[str | None, str, str | None, bool]:
    result = _run_isolated_pdf_worker(
        data,
        stop_after_chars=stop_after_chars,
    )
    status = result.get("status")
    error = result.get("error")
    error_text = error if isinstance(error, str) else None
    if status == "ok":
        text = result.get("text")
        text = text if isinstance(text, str) else None
        method = result.get("method")
        method = method if isinstance(method, str) else "pdf_extraction_failed"
        resource_truncated = result.get("resource_truncated") is True
        if text is not None and len(text) > stop_after_chars:
            text = text[:stop_after_chars]
            resource_truncated = True
            error_text = "pdf_projection_limit_reached"
        return text, method, error_text, resource_truncated
    if status == "resource_limit":
        return None, "pdf_resource_limit_exceeded", error_text, True
    if status == "isolation_unsupported":
        return None, "pdf_text_projection_unsupported", error_text, False
    if status == "parser_error":
        return None, "pdf_extraction_failed", error_text, False
    return None, "pdf_extraction_failed", error_text, False


def _extract_docx(
    data: bytes, *, stop_after_chars: int
) -> tuple[str | None, str, str | None, bool]:
    try:
        _preflight_ooxml_archive(data, document_kind="docx")
        from docx import Document  # type: ignore[import-not-found]

        document = Document(io.BytesIO(data))
        parts: list[str] = []
        length = 0
        resource_truncated = False
        for paragraph_number, paragraph in enumerate(document.paragraphs, 1):
            if paragraph_number > MAX_DOCX_PARAGRAPHS:
                resource_truncated = True
                break
            value = paragraph.text.strip()
            if value:
                length, char_limit_reached = _append_bounded_text(
                    parts,
                    value,
                    length=length,
                    stop_after_chars=stop_after_chars,
                )
                if char_limit_reached:
                    resource_truncated = True
                    break
            if length >= stop_after_chars:
                break
        row_count = 0
        cell_count = 0
        if length < stop_after_chars and not resource_truncated:
            for table_number, table in enumerate(document.tables, 1):
                if table_number > MAX_DOCX_TABLES:
                    resource_truncated = True
                    break
                for row in table.rows:
                    row_count += 1
                    if row_count > MAX_DOCX_ROWS:
                        resource_truncated = True
                        break
                    cell_count += len(row.cells)
                    if cell_count > MAX_DOCX_CELLS:
                        resource_truncated = True
                        break
                    value = "\t".join(cell.text.strip() for cell in row.cells).strip()
                    if value:
                        length, char_limit_reached = _append_bounded_text(
                            parts,
                            value,
                            length=length,
                            stop_after_chars=stop_after_chars,
                        )
                        if char_limit_reached:
                            resource_truncated = True
                            break
                    if length >= stop_after_chars:
                        break
                if resource_truncated or length >= stop_after_chars:
                    break
        text = "\n".join(parts).strip()
        error = "docx_projection_limit_reached" if resource_truncated else None
        return text or None, "python_docx", error, resource_truncated
    except _ProjectionResourceLimit as exc:
        return None, "docx_resource_limit_exceeded", str(exc), True
    except Exception as exc:  # noqa: BLE001 - third-party parser boundary
        return None, "python_docx_failed", type(exc).__name__, False


def _extract_pptx(
    data: bytes, *, stop_after_chars: int
) -> tuple[str | None, str, str | None, bool]:
    try:
        _preflight_ooxml_archive(data, document_kind="pptx")
        from pptx import Presentation  # type: ignore[import-not-found]

        presentation = Presentation(io.BytesIO(data))
        parts: list[str] = []
        length = 0
        shape_count = 0
        resource_truncated = False
        for slide_number, slide in enumerate(presentation.slides, 1):
            if slide_number > MAX_PPTX_SLIDES:
                resource_truncated = True
                break
            slide_parts: list[str] = []
            for shape in slide.shapes:
                shape_count += 1
                if shape_count > MAX_PPTX_SHAPES:
                    resource_truncated = True
                    break
                value = getattr(shape, "text", None)
                if isinstance(value, str) and value.strip():
                    slide_parts.append(value.strip())
            if slide_parts:
                value = f"Slide {slide_number}\n" + "\n".join(slide_parts)
                length, char_limit_reached = _append_bounded_text(
                    parts,
                    value,
                    length=length,
                    stop_after_chars=stop_after_chars,
                )
                if char_limit_reached:
                    resource_truncated = True
            if resource_truncated or length >= stop_after_chars:
                break
        text = "\n\n".join(parts).strip()
        error = "pptx_projection_limit_reached" if resource_truncated else None
        return text or None, "python_pptx", error, resource_truncated
    except _ProjectionResourceLimit as exc:
        return None, "pptx_resource_limit_exceeded", str(exc), True
    except Exception as exc:  # noqa: BLE001 - third-party parser boundary
        return None, "python_pptx_failed", type(exc).__name__, False


def _extract_xlsx(
    data: bytes, *, stop_after_chars: int
) -> tuple[str | None, str, str | None, bool]:
    try:
        _preflight_ooxml_archive(data, document_kind="xlsx")
        from openpyxl import load_workbook  # type: ignore[import-not-found]

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        length = 0
        cell_count = 0
        resource_truncated = False
        stop_iteration = False
        try:
            worksheets = workbook.worksheets
            if len(worksheets) > MAX_XLSX_WORKSHEETS:
                resource_truncated = True
            for worksheet in worksheets[:MAX_XLSX_WORKSHEETS]:
                heading = f"Sheet: {worksheet.title}"
                length, char_limit_reached = _append_bounded_text(
                    parts,
                    heading,
                    length=length,
                    stop_after_chars=stop_after_chars,
                )
                if char_limit_reached:
                    resource_truncated = True
                    stop_iteration = True
                    break
                if worksheet.max_row > MAX_XLSX_ROWS:
                    resource_truncated = True
                if worksheet.max_column > MAX_XLSX_COLUMNS:
                    resource_truncated = True
                bounded_max_row = min(worksheet.max_row, MAX_XLSX_ROWS)
                bounded_max_column = min(worksheet.max_column, MAX_XLSX_COLUMNS)
                for row_number, row in enumerate(
                    worksheet.iter_rows(
                        values_only=True,
                        max_row=bounded_max_row,
                        max_col=bounded_max_column,
                    ),
                    1,
                ):
                    if row_number > MAX_XLSX_ROWS:
                        resource_truncated = True
                        stop_iteration = True
                        break
                    cell_count += len(row)
                    if cell_count > MAX_XLSX_CELLS:
                        resource_truncated = True
                        stop_iteration = True
                        break
                    value = "\t".join(
                        "" if cell is None else str(cell) for cell in row
                    ).rstrip()
                    if value:
                        length, char_limit_reached = _append_bounded_text(
                            parts,
                            value,
                            length=length,
                            stop_after_chars=stop_after_chars,
                        )
                        if char_limit_reached:
                            resource_truncated = True
                            stop_iteration = True
                            break
                    if length >= stop_after_chars:
                        stop_iteration = True
                        break
                if stop_iteration:
                    break
        finally:
            workbook.close()
        text = "\n".join(parts).strip()
        error = "xlsx_projection_limit_reached" if resource_truncated else None
        return text or None, "openpyxl_read_only", error, resource_truncated
    except _ProjectionResourceLimit as exc:
        return None, "xlsx_resource_limit_exceeded", str(exc), True
    except Exception as exc:  # noqa: BLE001 - third-party parser boundary
        return None, "openpyxl_failed", type(exc).__name__, False


def _extract_html(data: bytes) -> tuple[str | None, str, str | None]:
    try:
        from bs4 import BeautifulSoup  # type: ignore[import-not-found]

        decoded, _encoding, _source_truncated = _decode_text(data)
        soup = BeautifulSoup(decoded, "html.parser")
        for element in soup(("script", "style")):
            element.decompose()
        return soup.get_text(separator="\n", strip=True) or None, "beautifulsoup", None
    except Exception as exc:  # noqa: BLE001 - optional parser fallback
        decoded, _encoding, _source_truncated = _decode_text(data)
        fallback = re.sub(r"<[^>]+>", " ", decoded)
        fallback = re.sub(r"\s+", " ", unescape(fallback)).strip()
        return fallback or None, "html_stdlib_fallback", type(exc).__name__


def _extract_eml(data: bytes) -> tuple[str | None, str, str | None]:
    try:
        message = BytesParser(policy=email_policy.default).parsebytes(data)
        parts = [
            f"{name}: {message.get(name)}"
            for name in ("From", "To", "Subject", "Date")
            if message.get(name)
        ]
        body = message.get_body(preferencelist=("plain",))
        if body is not None:
            value = body.get_content()
            if isinstance(value, str) and value.strip():
                parts.extend(("", value.strip()))
        text = "\n".join(parts).strip()
        return text or None, "email_stdlib", None
    except Exception as exc:  # noqa: BLE001 - stdlib parser boundary
        return None, "email_stdlib_failed", type(exc).__name__


def extract_file_bytes_text_projection(
    *,
    data: bytes,
    content_type: str | None,
    original_filename: str | None,
    max_text_chars: int = DEFAULT_MAX_TEXT_CHARS,
) -> dict[str, Any]:
    """Return a bounded text projection without creating durable state."""

    if not isinstance(data, (bytes, bytearray)):
        raise TypeError("data must be bytes")
    if isinstance(max_text_chars, bool) or not isinstance(max_text_chars, int):
        raise TypeError("max_text_chars must be an integer")
    if max_text_chars <= 0 or max_text_chars > MAX_TEXT_CHARS:
        raise ValueError(f"max_text_chars must be between 1 and {MAX_TEXT_CHARS}")

    if len(data) > MAX_INPUT_BYTES:
        return {
            "text_length": 0,
            "text_truncated": True,
            "text_extraction": "input_resource_limit_exceeded",
            "text_extraction_error": "input_size_limit_exceeded",
        }

    data_bytes = bytes(data)
    content_type_token = _content_type_token(content_type)
    filename_token = _filename_token(original_filename)
    stop_after_chars = max_text_chars + 1
    resource_truncated = False
    source_truncated = False

    if content_type_token == "application/pdf" or filename_token.endswith(".pdf"):
        text, method, error, resource_truncated = _extract_pdf(
            data_bytes,
            stop_after_chars=stop_after_chars,
        )
        encoding = "utf-8" if text else None
    elif "wordprocessingml.document" in content_type_token or filename_token.endswith(
        ".docx"
    ):
        text, method, error, resource_truncated = _extract_docx(
            data_bytes,
            stop_after_chars=stop_after_chars,
        )
        encoding = "utf-8" if text else None
    elif "presentationml.presentation" in content_type_token or filename_token.endswith(
        ".pptx"
    ):
        text, method, error, resource_truncated = _extract_pptx(
            data_bytes,
            stop_after_chars=stop_after_chars,
        )
        encoding = "utf-8" if text else None
    elif "spreadsheetml.sheet" in content_type_token or filename_token.endswith(
        ".xlsx"
    ):
        text, method, error, resource_truncated = _extract_xlsx(
            data_bytes,
            stop_after_chars=stop_after_chars,
        )
        encoding = "utf-8" if text else None
    elif content_type_token in {
        "text/html",
        "application/xhtml+xml",
    } or filename_token.endswith((".html", ".htm", ".xhtml")):
        if len(data_bytes) > MAX_MARKUP_OR_EMAIL_BYTES:
            text = None
            method = "html_resource_limit_exceeded"
            error = "markup_size_limit_exceeded"
            encoding = None
            resource_truncated = True
        else:
            text, method, error = _extract_html(data_bytes)
            encoding = "utf-8" if text else None
    elif content_type_token == "message/rfc822" or filename_token.endswith(".eml"):
        if len(data_bytes) > MAX_MARKUP_OR_EMAIL_BYTES:
            text = None
            method = "email_resource_limit_exceeded"
            error = "email_size_limit_exceeded"
            encoding = None
            resource_truncated = True
        else:
            text, method, error = _extract_eml(data_bytes)
            encoding = "utf-8" if text else None
    elif content_type_token.startswith("image/") or filename_token.endswith(
        (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif")
    ):
        # OCR engines may render arbitrarily large raster surfaces and spawn
        # external processes.  This pure projection has no dependable way to
        # confine either, so image OCR belongs in an explicitly bounded
        # capability rather than an implicit fallback here.
        text = None
        method = "image_text_projection_unsupported"
        error = None
        encoding = None
    elif (
        content_type_token.startswith("text/")
        or content_type_token
        in {
            "application/json",
            "application/ld+json",
            "application/xml",
            "application/yaml",
            "application/x-yaml",
        }
        or filename_token.endswith(
            (
                ".txt",
                ".md",
                ".markdown",
                ".csv",
                ".tsv",
                ".json",
                ".xml",
                ".yaml",
                ".yml",
                ".tex",
            )
        )
        or _looks_textual(data_bytes)
    ):
        text, encoding, source_truncated = _decode_text(
            data_bytes,
            stop_after_chars=stop_after_chars,
        )
        method = "text_decode"
        error = None
    else:
        text = None
        method = "unsupported_binary"
        error = None
        encoding = None

    text = text.strip() if isinstance(text, str) else None
    original_length = len(text) if text is not None else 0
    truncated = (
        resource_truncated or source_truncated or original_length > max_text_chars
    )
    projected_text = text[:max_text_chars] if text is not None else None

    payload: dict[str, Any] = {
        "text_length": len(projected_text) if projected_text is not None else 0,
        "text_truncated": truncated,
        "text_extraction": method,
    }
    if projected_text is not None:
        payload["text"] = projected_text
    if encoding:
        payload["encoding"] = encoding
    if error:
        payload["text_extraction_error"] = error[:240]
    return payload


__all__ = [
    "DEFAULT_MAX_TEXT_CHARS",
    "MAX_INPUT_BYTES",
    "MAX_TEXT_CHARS",
    "extract_file_bytes_text_projection",
]
