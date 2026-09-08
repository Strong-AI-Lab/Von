"""Source-preserving slide evidence; interpretation policy belongs to the caller.

Renderings are derived cache entries, keyed by source bytes and renderer version.
Every request still checks access to the source file before consulting the cache.
"""

from __future__ import annotations

import hashlib
import io
import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from .file_bytes_text_projection_service import (
    MAX_INPUT_BYTES,
    _preflight_ooxml_archive,
    detect_ooxml_content_type,
)

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
EVIDENCE_VERSION = "presentation_evidence.v1"


def native_slides(data: bytes) -> list[dict[str, Any]]:
    from pptx import Presentation

    _preflight_ooxml_archive(data, document_kind="pptx")
    presentation = Presentation(io.BytesIO(data))
    slides = []

    def shape_text(shapes):
        parts = []
        for shape in shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                parts.extend(
                    "\t".join(cell.text for cell in row.cells)
                    for row in shape.table.rows
                )
            if hasattr(shape, "shapes"):
                parts.extend(shape_text(shape.shapes))
        return parts

    for number, slide in enumerate(presentation.slides, 1):
        notes_frame = (
            slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        )
        notes = notes_frame.text if notes_frame is not None else ""
        slides.append(
            {
                "slide_number": number,
                "text": "\n".join(shape_text(slide.shapes)),
                "notes": notes,
            }
        )
    return slides


def render_pdf(data: bytes, *, source_sha256: str) -> tuple[bytes, str, bool]:
    from .blob_store import get_blob_store_from_env

    executable = shutil.which("soffice")
    if not executable:
        raise ValueError(
            "presentation_renderer_unavailable: install LibreOffice (soffice)"
        )
    version = (
        subprocess.run(
            [executable, "--version"], capture_output=True, check=True, timeout=15
        )
        .stdout.decode()
        .strip()
    )
    renderer = f"libreoffice:{version}:{EVIDENCE_VERSION}"
    cache_hash = hashlib.sha256(f"{source_sha256}:{renderer}".encode()).hexdigest()
    key = f"derived/presentations/{cache_hash}.pdf"
    store = get_blob_store_from_env()
    try:
        cached = store.get_bytes(key)
    except Exception:  # noqa: BLE001 - optional derived cache; render from authorised source
        cached = None
    if cached and cached.startswith(b"%PDF-"):
        return cached, renderer, True
    _preflight_ooxml_archive(data, document_kind="pptx")
    with tempfile.TemporaryDirectory(prefix="von-slides-") as root:
        work = Path(root)
        (work / "deck.pptx").write_bytes(data)
        profile = (work / "profile").as_uri()
        subprocess.run(
            [
                executable,
                f"-env:UserInstallation={profile}",
                "--headless",
                "--convert-to",
                "pdf:impress_pdf_Export",
                "--outdir",
                str(work),
                str(work / "deck.pptx"),
            ],
            check=True,
            capture_output=True,
            timeout=90,
        )
        pdf_path = work / "deck.pdf"
        if not pdf_path.exists() or pdf_path.stat().st_size > MAX_INPUT_BYTES:
            raise ValueError("presentation_render_failed_or_too_large")
        pdf = pdf_path.read_bytes()
    if not pdf.startswith(b"%PDF-"):
        raise ValueError("presentation_render_invalid_pdf")
    store.put_bytes(key, pdf, content_type="application/pdf")
    return pdf, renderer, False


def read_presentation_slides(
    *,
    file_copy_concept_id: str,
    user_concept_id: str,
    organisation_concept_id: str | None = None,
    offset: int = 0,
    limit: int = 4,
    include_ocr: bool = False,
) -> dict[str, Any]:
    import fitz

    from .computer_file_copy_service import fetch_file_copy_bytes
    from .conversation_image_service import store_image

    if not user_concept_id:
        raise PermissionError("Authenticated source access is required")
    if isinstance(offset, bool) or not isinstance(offset, int) or offset < 0:
        raise ValueError("offset must be a non-negative integer")
    if isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= 4:
        raise ValueError("limit must be 1 to 4 slides")
    source = fetch_file_copy_bytes(
        file_copy_concept_id=file_copy_concept_id,
        user_concept_id=user_concept_id,
        organisation_concept_id=organisation_concept_id,
        max_bytes=MAX_INPUT_BYTES,
    )
    if not source.get("success"):
        raise PermissionError("Presentation source unavailable to this actor")
    data = source["data"]
    source_hash = hashlib.sha256(data).hexdigest()
    if detect_ooxml_content_type(data) == PPTX_MIME:
        native = native_slides(data)
        pdf, renderer, cache_hit = render_pdf(data, source_sha256=source_hash)
    elif data.startswith(b"%PDF-"):
        native = None
        pdf, renderer, cache_hit = data, f"pymupdf:{fitz.VersionBind}", False
    else:
        raise ValueError("presentation_format_unsupported: use PPTX or PDF")
    with fitz.open(stream=pdf, filetype="pdf") as document:
        count = len(document)
        if count > 500:
            raise ValueError("presentation_slide_limit_exceeded")
        if native is not None and len(native) != count:
            raise ValueError("presentation_render_slide_count_mismatch")
        results, attachments = [], []
        for page_index in range(offset, min(offset + limit, count)):
            page = document[page_index]
            scale = min(2, 1800 / max(page.rect.width, page.rect.height))
            png = page.get_pixmap(
                matrix=fitz.Matrix(scale, scale), alpha=False
            ).tobytes("png")
            provenance = {
                "kind": "presentation_slide",
                "source_file_copy_concept_id": file_copy_concept_id,
                "source_sha256": source_hash,
                "slide_number": page_index + 1,
                "renderer": renderer,
                "rasteriser": f"pymupdf:{fitz.VersionBind}",
                "evidence_version": EVIDENCE_VERSION,
            }
            image = store_image(
                data=png,
                filename=f"slide-{page_index + 1}.png",
                user_concept_id=user_concept_id,
                provenance=provenance,
            )
            slide = (
                dict(native[page_index])
                if native
                else {
                    "slide_number": page_index + 1,
                    "text": page.get_text(),
                    "notes": None,
                }
            )
            slide["text_truncated"] = len(slide["text"]) > 20000
            slide["notes_truncated"] = len(slide.get("notes") or "") > 20000
            slide["text"] = slide["text"][:20000]
            if slide.get("notes"):
                slide["notes"] = slide["notes"][:20000]
            if include_ocr:
                slide["ocr"] = cached_ocr(png)
            slide["image"] = image
            results.append(slide)
            attachments.append(image)
    return {
        "success": True,
        "schema_version": EVIDENCE_VERSION,
        "source_file_copy_concept_id": file_copy_concept_id,
        "source_sha256": source_hash,
        "slide_count": count,
        "offset": offset,
        "next_offset": offset + len(results) if offset + len(results) < count else None,
        "renderer": renderer,
        "render_cache_hit": cache_hit,
        "slides": results,
        "image_attachments": attachments,
        "note": "Slide images are renderings; native text/notes and visual interpretation are distinct evidence. Cite slide numbers and source image URLs. Presenter is not necessarily paper author.",
    }


def cached_ocr(png: bytes) -> dict[str, Any]:
    """Cache an extraction independently of later model interpretations."""
    import pytesseract
    from PIL import Image

    from .blob_store import get_blob_store_from_env

    engine = (
        f"tesseract:{pytesseract.get_tesseract_version()}:default:{EVIDENCE_VERSION}"
    )
    source_hash = hashlib.sha256(png).hexdigest()
    cache_hash = hashlib.sha256(f"{source_hash}:{engine}".encode()).hexdigest()
    store = get_blob_store_from_env()
    key = f"derived/presentation-ocr/{cache_hash}.json"
    try:
        cached = json.loads(store.get_bytes(key))
        if cached.get("image_sha256") == source_hash and cached.get("engine") == engine:
            return {**cached, "cache_hit": True}
    except Exception:  # noqa: BLE001 - optional derived cache
        cached = None
    try:
        with Image.open(io.BytesIO(png)) as image:
            text = pytesseract.image_to_string(image, timeout=30)
        result = {
            "image_sha256": source_hash,
            "engine": engine,
            "text": text[:20000],
            "text_truncated": len(text) > 20000,
            "cache_hit": False,
        }
        store.put_bytes(
            key, json.dumps(result).encode(), content_type="application/json"
        )
        return result
    except (RuntimeError, OSError) as exc:
        return {
            "image_sha256": source_hash,
            "engine": engine,
            "error": type(exc).__name__,
        }
