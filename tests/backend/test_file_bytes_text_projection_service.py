from __future__ import annotations

import io
import subprocess
import zipfile

import pytest


def _project(
    data: bytes,
    *,
    content_type: str,
    filename: str,
    max_text_chars: int = 20_000,
) -> dict:
    from src.backend.services.file_bytes_text_projection_service import (
        extract_file_bytes_text_projection,
    )

    return extract_file_bytes_text_projection(
        data=data,
        content_type=content_type,
        original_filename=filename,
        max_text_chars=max_text_chars,
    )


def _zip_bytes(*members: tuple[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in members:
            archive.writestr(name, value)
    return buffer.getvalue()


def test_source_size_boundary_prevents_parser_dispatch(monkeypatch):
    from src.backend.services import file_bytes_text_projection_service as service

    monkeypatch.setattr(service, "MAX_INPUT_BYTES", 4)
    monkeypatch.setattr(
        service,
        "_looks_textual",
        lambda _data: (_ for _ in ()).throw(
            AssertionError("oversized input must not reach format detection")
        ),
    )

    result = _project(
        b"12345",
        content_type="application/octet-stream",
        filename="source.bin",
    )

    assert result == {
        "text_length": 0,
        "text_truncated": True,
        "text_extraction": "input_resource_limit_exceeded",
        "text_extraction_error": "input_size_limit_exceeded",
    }


def test_plain_text_decodes_only_a_bounded_prefix():
    result = _project(
        b"a" * 1_000_000,
        content_type="text/plain",
        filename="large.txt",
        max_text_chars=11,
    )

    assert result["text"] == "a" * 11
    assert result["text_length"] == 11
    assert result["text_truncated"] is True
    assert result["text_extraction"] == "text_decode"


def test_ooxml_member_count_is_rejected_before_document_parser(monkeypatch):
    from src.backend.services import file_bytes_text_projection_service as service

    monkeypatch.setattr(service, "MAX_ARCHIVE_MEMBERS", 2)
    source = _zip_bytes(
        ("[Content_Types].xml", b"content"),
        ("word/document.xml", b"document"),
        ("word/styles.xml", b"styles"),
    )

    result = _project(
        source,
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        filename="bounded.docx",
    )

    assert result["text_extraction"] == "docx_resource_limit_exceeded"
    assert result["text_extraction_error"] == "archive_member_limit_exceeded"
    assert result["text_length"] == 0
    assert result["text_truncated"] is True


def test_ooxml_expansion_is_rejected_before_document_parser(monkeypatch):
    from src.backend.services import file_bytes_text_projection_service as service

    monkeypatch.setattr(service, "MAX_ARCHIVE_MEMBER_BYTES", 100)
    monkeypatch.setattr(service, "MAX_ARCHIVE_EXPANDED_BYTES", 30)
    source = _zip_bytes(
        ("[Content_Types].xml", b"a" * 20),
        ("word/document.xml", b"b" * 20),
    )

    result = _project(
        source,
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        filename="expanded.docx",
    )

    assert result["text_extraction"] == "docx_resource_limit_exceeded"
    assert result["text_extraction_error"] == "archive_expansion_limit_exceeded"


def test_pdf_page_projection_stops_at_local_page_boundary(monkeypatch):
    from src.backend.services import file_bytes_text_projection_service as service

    if not service._pdf_parser_isolation_supported():
        pytest.skip("hard PDF parser isolation is not available on this host")
    import fitz

    document = fitz.open()
    for page_number in range(1, 4):
        page = document.new_page()
        page.insert_text((72, 72), f"page-{page_number}-content")
    source = document.tobytes()
    document.close()
    monkeypatch.setattr(service, "MAX_PDF_PAGES", 2)

    worker_result = service._run_isolated_pdf_worker(
        source,
        stop_after_chars=service.MAX_TEXT_CHARS,
    )
    if worker_result["status"] != "ok":
        direct_worker = subprocess.run(
            service._pdf_worker_command(stop_after_chars=service.MAX_TEXT_CHARS),
            input=source,
            capture_output=True,
            check=False,
        )
        pytest.fail(
            repr(
                {
                    "worker_result": worker_result,
                    "direct_return_code": direct_worker.returncode,
                    "direct_stdout": direct_worker.stdout[:1_000],
                    "direct_stderr": direct_worker.stderr[:1_000],
                }
            )
        )

    result = _project(
        source,
        content_type="application/pdf",
        filename="many-pages.pdf",
    )

    assert result["text_extraction"] == "pymupdf", result
    assert "page-1-content" in result["text"]
    assert "page-2-content" in result["text"]
    assert "page-3-content" not in result["text"]
    assert result["text_truncated"] is True
    assert result["text_extraction_error"] == "pdf_projection_limit_reached"


def test_pdf_failure_has_no_unbounded_secondary_parser():
    from src.backend.services import file_bytes_text_projection_service as service

    if not service._pdf_parser_isolation_supported():
        pytest.skip("hard PDF parser isolation is not available on this host")

    result = _project(
        b"%PDF-broken",
        content_type="application/pdf",
        filename="broken.pdf",
    )

    assert result["text_extraction"] == "pdf_extraction_failed"
    assert isinstance(result["text_extraction_error"], str)


def test_docx_paragraph_boundary_is_checked_before_document_parser(monkeypatch):
    import docx
    from docx import Document

    from src.backend.services import file_bytes_text_projection_service as service

    document = Document()
    document.add_paragraph("first paragraph")
    document.add_paragraph("second paragraph")
    document.add_paragraph("third paragraph")
    buffer = io.BytesIO()
    document.save(buffer)
    monkeypatch.setattr(service, "MAX_DOCX_PARAGRAPHS", 2)
    monkeypatch.setattr(
        docx,
        "Document",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("over-complex DOCX must not reach python-docx")
        ),
    )

    result = _project(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        filename="paragraphs.docx",
    )

    assert result["text_extraction"] == "docx_resource_limit_exceeded"
    assert result["text_extraction_error"] == "docx_paragraph_limit_exceeded"
    assert "text" not in result


def test_pptx_shape_boundary_is_checked_before_presentation_parser(monkeypatch):
    import pptx
    from pptx import Presentation
    from pptx.util import Inches

    from src.backend.services import file_bytes_text_projection_service as service

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index in range(1, 3):
        box = slide.shapes.add_textbox(Inches(1), Inches(index), Inches(4), Inches(1))
        box.text = f"shape-{index}-content"
    buffer = io.BytesIO()
    presentation.save(buffer)
    monkeypatch.setattr(service, "MAX_PPTX_SHAPES", 1)
    monkeypatch.setattr(
        pptx,
        "Presentation",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("over-complex PPTX must not reach python-pptx")
        ),
    )

    result = _project(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
        filename="shapes.pptx",
    )

    assert result["text_extraction"] == "pptx_resource_limit_exceeded"
    assert result["text_extraction_error"] == "pptx_shape_limit_exceeded"
    assert "text" not in result


def test_xlsx_row_boundary_is_checked_before_workbook_parser(monkeypatch):
    import openpyxl
    from openpyxl import Workbook

    from src.backend.services import file_bytes_text_projection_service as service

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(("row-1", "value-1"))
    worksheet.append(("row-2", "value-2"))
    worksheet.append(("row-3", "value-3"))
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    monkeypatch.setattr(service, "MAX_XLSX_ROWS", 2)
    monkeypatch.setattr(service, "MAX_XLSX_COLUMNS", 2)
    monkeypatch.setattr(
        openpyxl,
        "load_workbook",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("over-complex XLSX must not reach openpyxl")
        ),
    )

    result = _project(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        filename="rows.xlsx",
    )

    assert result["text_extraction"] == "xlsx_resource_limit_exceeded"
    assert result["text_extraction_error"] == "xlsx_row_limit_exceeded"
    assert "text" not in result


def test_images_do_not_implicitly_invoke_unbounded_ocr(monkeypatch):
    from src.backend.services import file_copy_interpretation_service

    monkeypatch.setattr(
        file_copy_interpretation_service,
        "extract_image_ocr_text",
        lambda _data: (_ for _ in ()).throw(
            AssertionError("ephemeral projection must not invoke OCR")
        ),
    )

    result = _project(
        b"image bytes",
        content_type="image/png",
        filename="scan.png",
    )

    assert result == {
        "text_length": 0,
        "text_truncated": False,
        "text_extraction": "image_text_projection_unsupported",
    }
