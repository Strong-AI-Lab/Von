import hashlib
import io

import fitz
import pytest
from pptx import Presentation

from src.backend.services import presentation_evidence_service as service


def deck_bytes():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Test framework"
    slide.notes_slide.notes_text_frame.text = "Check the arrow direction."
    slide.shapes.add_table(1, 2, 0, 0, 100000, 100000).table.cell(
        0, 1
    ).text = "Result 0.6671"
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def pdf_bytes():
    doc = fitz.open()
    for n in range(3):
        page = doc.new_page(width=600, height=400)
        page.insert_text((40, 40), f"Slide {n + 1}")
    return doc.tobytes()


def test_native_slide_notes_tables_and_order():
    slides = service.native_slides(deck_bytes())
    assert len(slides) == 1
    assert slides[0]["slide_number"] == 1
    assert "Test framework" in slides[0]["text"]
    assert "0.6671" in slides[0]["text"]
    assert slides[0]["notes"] == "Check the arrow direction."


def test_source_authority_before_render_or_cache(monkeypatch):
    monkeypatch.setattr(
        "src.backend.services.computer_file_copy_service.fetch_file_copy_bytes",
        lambda **kwargs: {"success": False},
    )
    monkeypatch.setattr(
        service, "render_pdf", lambda *args, **kwargs: pytest.fail("must not render")
    )
    with pytest.raises(PermissionError):
        service.read_presentation_slides(
            file_copy_concept_id="#V#private", user_concept_id="#V#other"
        )


def test_pdf_pagination_preserves_source_and_delivers_native_images(monkeypatch):
    data = pdf_bytes()
    monkeypatch.setattr(
        "src.backend.services.computer_file_copy_service.fetch_file_copy_bytes",
        lambda **kwargs: {"success": True, "data": data},
    )
    stored = []

    def store(**kwargs):
        stored.append(kwargs)
        return {
            "concept_id": f"#V#image_{len(stored)}",
            "provenance": kwargs["provenance"],
            "url": "/image",
        }

    monkeypatch.setattr(
        "src.backend.services.conversation_image_service.store_image", store
    )
    result = service.read_presentation_slides(
        file_copy_concept_id="#V#deck", user_concept_id="#V#owner", offset=1, limit=1
    )
    assert result["slide_count"] == 3
    assert result["next_offset"] == 2
    assert result["slides"][0]["slide_number"] == 2
    assert "Slide 2" in result["slides"][0]["text"]
    assert result["source_sha256"] == hashlib.sha256(data).hexdigest()
    assert len(result["image_attachments"]) == 1
    assert stored[0]["data"].startswith(b"\x89PNG")
    assert stored[0]["user_concept_id"] == "#V#owner"
    assert stored[0]["provenance"]["source_file_copy_concept_id"] == "#V#deck"


def test_ocr_cache_invalidates_on_image_change(monkeypatch):
    import pytesseract
    from PIL import Image

    cache, calls = {}, []

    class Store:
        def get_bytes(self, key):
            return cache[key]

        def put_bytes(self, key, data, **kwargs):
            cache[key] = data

    monkeypatch.setattr(
        "src.backend.services.blob_store.get_blob_store_from_env", lambda: Store()
    )
    monkeypatch.setattr(pytesseract, "get_tesseract_version", lambda: "test-1")
    monkeypatch.setattr(
        pytesseract,
        "image_to_string",
        lambda *args, **kwargs: calls.append(1) or "x -> y",
    )

    def png(colour):
        out = io.BytesIO()
        Image.new("RGB", (10, 10), colour).save(out, "PNG")
        return out.getvalue()

    first = service.cached_ocr(png("red"))
    second = service.cached_ocr(png("red"))
    third = service.cached_ocr(png("blue"))
    assert not first["cache_hit"] and second["cache_hit"] and not third["cache_hit"]
    assert len(calls) == 2
    monkeypatch.setattr(pytesseract, "get_tesseract_version", lambda: "test-2")
    assert not service.cached_ocr(png("red"))["cache_hit"]
    assert len(calls) == 3


def test_notes_part_without_body_placeholder_is_valid():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    notes = slide.notes_slide
    frame = notes.notes_text_frame
    body = frame._txBody
    parent = body.getparent()
    parent.getparent().remove(parent)
    assert notes.notes_text_frame is None
    out = io.BytesIO()
    deck.save(out)
    assert service.native_slides(out.getvalue())[0]["notes"] == ""
